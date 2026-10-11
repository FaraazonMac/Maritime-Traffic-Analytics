"""Excel workbook: summary, hub congestion, live pivot, estimates driven by an Assumptions sheet, and the clean data."""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .clean import CLEAN_COLUMNS
from .estimates import SIZE_ORDER

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="0D1620")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="555555")
INPUT = Font(name=FONT, size=10, color="0000FF")
YELLOW = PatternFill("solid", fgColor="FFFF99")
LEVEL_FILL = {"green": PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE"),
              "yellow": PatternFill("solid", bgColor="FFEB9C", fgColor="FFEB9C"),
              "red": PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")}


def _py(v):
    if v is None or v is pd.NA:
        return None
    if isinstance(v, (float, np.floating)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v


def _table(ws, df, row=1, col=1, formats=None, widths=True):
    """Write df with a styled header. Returns the last row used."""
    formats = formats or {}
    for j, name in enumerate(df.columns):
        c = ws.cell(row=row, column=col + j, value=str(name))
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            c = ws.cell(row=row + i, column=col + j, value=_py(v))
            c.font = BODY
            fmt = formats.get(df.columns[j])
            if fmt:
                c.number_format = fmt
    if widths:
        for j, name in enumerate(df.columns):
            vals = [len(str(name))] + [len(str(_py(v))) for v in df.iloc[:200, j]]
            ws.column_dimensions[get_column_letter(col + j)].width = min(max(vals) + 2, 45)
    return row + len(df)


def _level_rules(ws, ref):
    for k, fill in LEVEL_FILL.items():
        ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=[f'"{k}"'], fill=fill))


def build(path, m, assump, df, stamp, region):
    wb = Workbook()
    col = {name: get_column_letter(i + 1) for i, name in enumerate(CLEAN_COLUMNS)}

    # ---------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = f"Maritime Traffic Analytics - {region}", TITLE
    ws["A2"], ws["A2"].font = f"Snapshot taken {m['kpis']['snapshot_utc']} UTC. Source: live AIS feed (AISStream), cleaned by run_all.py.", NOTE
    k = m["kpis"]
    rows = [
        ("Ships in the snapshot (after cleaning)", k["total_ships"], "#,##0"),
        ("Under way, speed 0.5 kn or more (%)", k["pct_moving"] / 100, "0.0%"), ("Idle: stationary but status says under way (%)", k["pct_idle"] / 100, "0.0%"),
        ("At anchor (%)", k["pct_at_anchor"] / 100, "0.0%"),
        ("At berth / moored (%)", k["pct_at_berth"] / 100, "0.0%"),
        ("Ships inside a port hub's radius", k["ships_near_hubs"], "#,##0"),
        ("Ships waiting at anchor near a hub", k["waiting_near_hubs"], "#,##0"),
        ("Dark-fleet flagged (%)", k["pct_dark"] / 100, "0.00%"),
        ("Vessel type unknown (%)", k["pct_type_unknown"] / 100, "0.0%"),
        ("With an IMO number (%)", k["pct_with_imo"] / 100, "0.0%"), ("IMO passes checksum (%)", k["pct_valid_imo"] / 100, "0.0%"),
        ("Destination usable (%)", k["pct_dest_usable"] / 100, "0.0%"),
        ("Destination resolved to a port after cleaning (%)", k["pct_dest_resolved"] / 100, "0.0%"),
        ("Destination matched by the app alone (%)", k["pct_dest_app_matched"] / 100, "0.0%"),
        ("Median speed of moving ships (kn)", k["median_speed_moving_kn"], "0.0"),
        ("Busiest hub (ships in radius)", k["busiest_hub"], None),
        ("Hub with most ships waiting at anchor", k["most_waiting_hub"], None),
    ]
    ws["A4"], ws["B4"], ws["C4"] = "Metric", "Value", "Note"
    for c in ("A4", "B4", "C4"):
        ws[c].font, ws[c].fill = HEAD_FONT, HEAD_FILL
    for i, (label, val, fmt) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=label).font = BODY
        c = ws.cell(row=i, column=2, value=_py(val))
        c.font = BODY
        if fmt:
            c.number_format = fmt
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value="Live check from the Clean_Data sheet (formula, recalculates):").font = BOLD
    ws.cell(row=r + 1, column=1, value="Ships in Clean_Data").font = BODY
    ws.cell(row=r + 1, column=2, value=f"=COUNTA(Clean_Data!{col['mmsi']}2:{col['mmsi']}{len(df) + 1})").font = BODY
    ws.cell(row=r + 2, column=1, value="Share moving").font = BODY
    c = ws.cell(row=r + 2, column=2, value=f'=COUNTIF(Clean_Data!{col["status_group"]}:{col["status_group"]},"moving")/B{r + 1}')
    c.font, c.number_format = BODY, "0.0%"
    r += 4
    ws.cell(row=r, column=1, value="Sheets").font = BOLD
    guide = [("Hub_Congestion", "Ships waiting at anchor near each hub, with green/yellow/red levels (relative to the other hubs)"),
             ("Fleet_Mix", "Live pivot (COUNTIFS on Clean_Data): vessel type x status"),
             ("Flags / Routes", "Top flag states; most common departure -> destination voyages"),
             ("Data_Quality", "Every cleaning rule with the number of rows it touched, and field completeness"),
             ("Assumptions", "Blue cells are inputs with sources. Change one and Estimates recalculates"),
             ("Estimates", "Fuel, CO2 and cost of ships waiting at anchor: ESTIMATES, low / central / high"),
             ("Clean_Data", "The cleaned dataset")]
    for i, (a, b) in enumerate(guide, start=r + 1):
        ws.cell(row=i, column=1, value=a).font = BODY
        ws.cell(row=i, column=2, value=b).font = BODY
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 40

    # ---------------- Hub congestion
    ws = wb.create_sheet("Hub_Congestion")
    ht = m["hub_table"][["hub", "zone", "ships_in_radius", "at_anchor", "idle", "at_berth", "moving", "waiting_broad", "avg_speed_moving_kn",
                         "pct_at_anchor", "level_waiting", "level_waiting_broad"]]
    last = _table(ws, ht, formats={"avg_speed_moving_kn": "0.0", "pct_at_anchor": "0.0"})
    _level_rules(ws, f"K2:L{last}")
    ws.cell(row=last + 2, column=1, value="Levels are relative: the lowest third of hubs is green, the middle third yellow, the top third red. "
                                          "They are not an industry standard.").font = NOTE
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "bar", "Ships waiting at anchor", 12, 16
    ch.add_data(Reference(ws, min_col=4, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ch.legend = None
    ws.add_chart(ch, "N2")
    ws.freeze_panes = "A2"

    # ---------------- Fleet mix (live pivot)
    ws = wb.create_sheet("Fleet_Mix")
    types = list(m["fleet_by_type"]["vessel_type"])
    ws["A1"], ws["A1"].font = "Vessel type x status (live COUNTIFS on Clean_Data)", BOLD
    heads = ["vessel_type", "at_anchor", "at_berth", "idle", "moving", "other", "total", "share"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    vt, sg = col["vessel_type"], col["status_group"]
    n_end = 3 + len(types)
    for i, t in enumerate(types, start=4):
        ws.cell(row=i, column=1, value=t).font = BODY
        for j in range(2, 7):
            c = ws.cell(row=i, column=j, value=f"=COUNTIFS(Clean_Data!${vt}:${vt},$A{i},Clean_Data!${sg}:${sg},{get_column_letter(j)}$3)")
            c.font = BODY
        ws.cell(row=i, column=7, value=f"=SUM(B{i}:F{i})").font = BODY
        c = ws.cell(row=i, column=8, value=f"=G{i}/G${n_end + 1}")
        c.font, c.number_format = BODY, "0.0%"
    ws.cell(row=n_end + 1, column=1, value="Total").font = BOLD
    for j in range(2, 8):
        L = get_column_letter(j)
        ws.cell(row=n_end + 1, column=j, value=f"=SUM({L}4:{L}{n_end})").font = BOLD
    ch = BarChart()
    ch.type, ch.title, ch.height, ch.width = "col", "Ships by vessel type", 9, 16
    ch.add_data(Reference(ws, min_col=7, min_row=3, max_row=n_end), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=n_end))
    ch.legend = None
    ws.add_chart(ch, "J3")
    r0 = n_end + 4
    ws.cell(row=r0 - 1, column=1, value="Type profile (computed in Python)").font = BOLD
    ft = m["fleet_by_type"][["vessel_type", "ships", "share_pct", "pct_moving", "avg_speed_moving_kn", "pct_dark"]]
    _table(ws, ft, row=r0, widths=False)
    ws.column_dimensions["A"].width = 18
    for L in "BCDEFGH":
        ws.column_dimensions[L].width = 14

    # ---------------- Flags, Routes
    ws = wb.create_sheet("Flags")
    _table(ws, m["flags_top"], formats={"share_pct": "0.0"})
    ws = wb.create_sheet("Segments")
    _table(ws, m["segments"])
    ws = wb.create_sheet("Routes")
    _table(ws, m["routes"], formats={"avg_speed_kn": "0.0"})
    ws.cell(row=len(m["routes"]) + 3, column=1,
            value="Departure ports are detected live (ship leaves a known port while the backend is watching); destinations are typed by captains.").font = NOTE

    # ---------------- Data quality
    ws = wb.create_sheet("Data_Quality")
    last = _table(ws, m["cleaning_log"], widths=False)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 40
    last = _table(ws, m["completeness"], row=last + 3, widths=False)

    # ---------------- Assumptions
    ws = wb.create_sheet("Assumptions")
    ws["A1"], ws["A1"].font = "Assumptions (blue = input you can change; yellow = key assumption)", BOLD
    ws["A2"], ws["A2"].font = "All values are published figures or clearly marked analyst assumptions. Nothing here is measured from AIS.", NOTE
    heads = ["parameter", "size_class", "low", "central", "high", "unit", "source", "note"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    arow = {}
    for i, rec in enumerate(assump.itertuples(index=False), start=5):
        arow[(rec.parameter, rec.size_class)] = i
        for j, name in enumerate(heads, start=1):
            c = ws.cell(row=i, column=j, value=_py(getattr(rec, name)))
            c.font = INPUT if name in ("low", "central", "high") else BODY
            if name in ("low", "central", "high"):
                c.fill = YELLOW
            if name in ("source", "note"):
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for L, w in zip("ABCDEFGH", (28, 22, 10, 10, 10, 18, 80, 60)):
        ws.column_dimensions[L].width = w

    # ---------------- Estimates (live formulas)
    ws = wb.create_sheet("Estimates")
    ws["A1"], ws["A1"].font = "ESTIMATE: daily fuel, CO2 and cost of ships waiting at anchor near a hub", TITLE
    ws["A2"], ws["A2"].font = ("These are estimates built from the Assumptions sheet, not measurements. Low uses every low assumption together, "
                               "high every high one, so the range is wide on purpose."), NOTE
    heads = ["Size class", "Ships waiting", "Fuel t/day low", "Fuel t/day central", "Fuel t/day high",
             "CO2 t/day low", "CO2 t/day central", "CO2 t/day high", "Cost USD/day low", "Cost USD/day central", "Cost USD/day high"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    price, co2, opex = (arow[("fuel_price_usd_per_t", "all")], arow[("co2_t_per_t_fuel", "all")], arow[("vessel_opex_usd_per_day", "all")])
    sc_col, hub_col = col["size_class"], col["hub"]
    for i, sc in enumerate(SIZE_ORDER, start=5):
        fr = arow[("fuel_at_anchor_t_per_day", sc)]
        ws.cell(row=i, column=1, value=sc).font = BODY
        ws.cell(row=i, column=2, value=(f'=COUNTIFS(Clean_Data!${sg}:${sg},"at_anchor",Clean_Data!${hub_col}:${hub_col},"<>",'
                                        f'Clean_Data!${sc_col}:${sc_col},$A{i})')).font = BODY
        for k, (L, sL) in enumerate((("C", "C"), ("D", "D"), ("E", "E"))):
            ws[f"{L}{i}"] = f"=$B{i}*Assumptions!${sL}${fr}"
        for k, (L, src, sL) in enumerate((("F", "C", "C"), ("G", "D", "D"), ("H", "E", "E"))):
            ws[f"{L}{i}"] = f"={src}{i}*Assumptions!${sL}${co2}"
        for L, sL in (("I", "C"), ("J", "D"), ("K", "E")):
            ws[f"{L}{i}"] = f"=$B{i}*(Assumptions!${sL}${opex}+Assumptions!${sL}${fr}*Assumptions!${sL}${price})"
        for L in "CDEFGHIJK":
            ws[f"{L}{i}"].font = BODY
            ws[f"{L}{i}"].number_format = "$#,##0" if L in "IJK" else "#,##0.0"
    t = 5 + len(SIZE_ORDER)
    ws.cell(row=t, column=1, value="Total").font = BOLD
    for L in "BCDEFGHIJK":
        c = ws[f"{L}{t}"]
        c.value = f"=SUM({L}5:{L}{t - 1})"
        c.font = BOLD
        c.number_format = "$#,##0" if L in "IJK" else ("#,##0" if L == "B" else "#,##0.0")
    notes = [
        "How to read: cost = ships x (vessel running cost per day + fuel per day x fuel price). Fuel = ships x fuel at anchor per day. CO2 = fuel x emission factor.",
        "This is the cost of ships sitting at anchor near a hub on a typical day. It is an upper bound for congestion-attributable cost, because some waiting is normal.",
        "A real congestion cost needs a baseline of normal waiting (see the time analysis once the history collector has run for several weeks).",
        "AIS contains no operator, cargo or schedule data, so no company-level loss can be computed.",
    ]
    for i, n in enumerate(notes, start=t + 2):
        ws.cell(row=i, column=1, value=n).font = NOTE
    ws.column_dimensions["A"].width = 24
    for L in "BCDEFGHIJK":
        ws.column_dimensions[L].width = 16
    ws.row_dimensions[4].height = 32

    # ---------------- Clean data
    ws = wb.create_sheet("Clean_Data")
    out = df[CLEAN_COLUMNS].astype(object).where(df[CLEAN_COLUMNS].notna(), None)
    ws.append(CLEAN_COLUMNS)
    for rec in out.itertuples(index=False):
        ws.append([_py(v) for v in rec])
    for j in range(1, len(CLEAN_COLUMNS) + 1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path
